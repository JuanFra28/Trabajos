"""
===============================================================================
 ESTANDARIZADOR DE DELAYS  ·  v1
===============================================================================
 Convierte el export crudo de delays (PU_Desc, Fecha_Inicio_Real, Turno,
 Reason_Level1..4, Delay, Maquina, Tipo) al formato estándar de la hoja
 Datos_Delays, que es el que usan reporte.py y el libro Planes_Delays_2026.xlsx:

   Maquina | Unidad | Fecha | Mes | Turno | Tipo | Area | Modulo | Cod |
   Familia | Sistema (L3) | Detalle (L4) | Horas | Minutos

 Qué hace:
   1. Lee el export (xlsx o csv) aunque cambien el nombre o el orden de columnas.
   2. Quita duplicados exactos y limpia textos (espacios, vacíos -> "(sin dato)").
   3. Máquina: valida el código (PF3, PI11...) y si falta lo saca de la unidad.
   4. Horas: usa Delay; si viene vacío lo calcula con Fin - Inicio.
   5. Tipo: agrupa variantes ("Electrical", "Electronico", "Avería Eléctrica"...)
      en Eléctricas / Mecánicas / ... y se queda solo con averías (configurable).
   6. Módulo: lo toma del nivel 2 del árbol de razones (Reason_Level2). Si el
      export no lo trae, lo recupera del histórico ya estandarizado.
   7. Área (agrupa módulos) y Familia (agrupa sistemas L3): tu histórico manda,
      y lo que sea nuevo se clasifica con reglas por palabras clave.
   8. L3/L4: unifica variantes de escritura ("Regulacion electrónica" =
      "Regulacion Electronica") y quita prefijos "Elec - " / "Mech - ".
   9. Escribe un Excel cuya primera hoja es Datos_Delays (lista para reporte.py),
      más hojas de resumen, calidad, pendientes por revisar y unificaciones.

 USO:
   python estandarizar_delays.py                          (usa ENTRADA)
   python estandarizar_delays.py delays.xlsx              (o pasa la entrada)
   python estandarizar_delays.py delays.xlsx salida.xlsx  (y la salida)

 Luego:
   python reporte.py Datos_Delays_estandar.xlsx

 Requiere: pandas numpy openpyxl
===============================================================================
"""

import difflib
import importlib.util
import os
import re
import subprocess
import sys
import unicodedata
import warnings
from collections import Counter
from datetime import datetime

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

# =============================================================================
# CONFIGURACION   <<<<<<  AJUSTA AQUI
# =============================================================================
ENTRADA = "delays.xlsx"                  # export crudo tal como te llega
HOJA_ENTRADA = None                      # None = primera hoja
SALIDA = "Datos_Delays_estandar.xlsx"    # primera hoja: Datos_Delays

# Histórico ya estandarizado (archivo, hoja). De aquí se recupera el módulo de
# los eventos que ya estaban clasificados y se aprende tu catálogo
# (módulo -> área, sistema L3 -> familia). Puede ser tu Planes_Delays_2026.xlsx
# o una salida anterior de este mismo script. Si no existe, se omite.
HISTORICO = [
    ("Planes_Delays_2026.xlsx", "Datos_Delays"),
]

# --- Filtros
TIPOS_INCLUIR = ["Eléctricas", "Mecánicas"]  # None = todos los tipos de parada
FECHA_DESDE = "2026-01-01"   # tu Datos_Delays arranca en 2026. None = todo
FECHA_HASTA = None           # "2026-08-31" para cortar en un mes cerrado

# --- Limpieza
QUITAR_DUPLICADOS = True     # filas idénticas en todas las columnas
QUITAR_PREFIJO_L4 = True     # "Mech - desgaste de faja" -> "Desgaste de faja"
PREFIJOS_L4 = ["Elec", "Mech", "Minor", "Menor", "Ex", "CG", "PP", "Man", "TNP"]
UNIFICAR_VARIANTES = True    # mismas palabras con distinta mayúscula/tilde/espacio
RECLASIFICAR_OTROS = True    # si tu histórico dejó un sistema en "Otros
                             # componentes" pero una regla lo reconoce, usa la regla
COLUMNAS_EXTRA = True        # agrega al final columnas de trazabilidad
                             # (Fin, Año, Origen módulo, textos originales...)

# --- Al terminar, generar el reporte HTML con reporte.py (misma carpeta)
EJECUTAR_REPORTE = False
REPORTE_PY = "reporte.py"

# --- Correcciones manuales (ganan a todo lo demás)
FAMILIAS_MANUAL = {
    # "Falla Neumatica": "Sistema Neumático / Aire",
}
MODULOS_MANUAL = {
    # "COD": ("Nombre del módulo", "Área"),
    # "AS4": ("Adhesivo De Laminado", "Adhesivos / Melters"),
}
ALIAS_MODULOS = {
    # texto del nivel 2 que no trae código -> código
    # "Adhesivo de construccion": "AS1",
}
TEXTOS_MANUAL = {
    # sinónimos en L3 / L4: "texto como viene": "texto estándar"
    # "control de web / guia": "Control de Web / Guia Fife",
}


# =============================================================================
# 1. CATALOGOS BASE (aprendidos de Datos_Delays; el histórico los actualiza)
# =============================================================================
SIN = "(sin dato)"
COLUMNAS = ["Maquina", "Unidad", "Fecha", "Mes", "Turno", "Tipo", "Area", "Modulo",
            "Cod", "Familia", "Sistema (L3)", "Detalle (L4)", "Horas", "Minutos"]

MODULOS_BASE = {
    "Adhesivos / Melters": {
        "AS1": "Adhesivo De Construccion", "AS2": "Adhesivo De Elasticos - Barrera",
        "AS3": "Adhesivo De Elastico", "MLT1": "Melter 1", "MLT2": "Melter 2",
    },
    "Aplicadores / Formación de producto": {
        "ALA": "Airlaid Application", "ARLD": "Airlaid Module",
        "BEA": "Back Ear Applicator", "BET": "Back Ear Forming",
        "BSA": "Backsheet Application", "CDR": "Rodillo De Contruccion",
        "CFA": "Formador De Barreras", "CFF": "Containment Flap Forming",
        "DBK": "Debulker", "EMB": "Embossing", "FBZ": "Fiberizer",
        "FEA": "Front Ear Applicator", "FLD": "Ear Folding", "FLF": "Fluff Forming",
        "FORM": "Forming Module", "FPA": "Frontal Patch Applicator",
        "LEA": "Leg Elastic", "MILL": "Mill Module", "NBM": "New Borning Module",
        "PAA": "Perfume Applicator", "PCM": "Pouch Cut Module", "PDB": "Predebulker",
        "PMS": "Pulp & Mill System", "RPA": "Release Paper Application",
        "SAM": "Super Absorbent", "SLA": "Surge Layer Applicator",
        "SLE": "Surge Layer Embosser", "SW": "Single Wrap",
        "SWA": "Solo Wrap Applicator", "UST": "Ultrasonic System Side Panel",
        "WEA": "Aplicador De Waist Elastic", "WFM": "Wing Folding Module",
    },
    "Corte y sellado": {
        "CPLN": "Coupler Module", "FCO": "Final Cut Off",
        "FCU": "Unidad De Corte Barreras", "FDC": "Corte Anatomico",
        "FNCT": "Final Cutting Module", "PCO": "Pad Cut Off",
        "PU03": "Unidad De Prensa 03", "QWM": "Quick Wrap Module",
        "RPT": "Repitch Module", "RTM": "Rotator Module", "SLCT": "Seal Cut Module",
        "SLNG": "Sealing Module", "TFLD": "Trifolder Module",
    },
    "Debobinadores / Unwinds": {
        "BEU": "Debobinador De Orejas Post", "BGFM": "Bag Film Unwinder",
        "BLU": "Barrier Layer Unwind", "BSU": "Backsheet Unwind",
        "CFU": "Debobinador De Barreras", "ETU": "Debobinador De Etiqueta",
        "FEU": "Front Ear Unwind", "ICU": "Inner Cover Unwind",
        "LEB": "Debobinador Elastico Barreras", "LEU": "Debobinador Elastico Pierna",
        "LNU": "Liner Unwind", "MFU": "Mechanical Fastener Unwind",
        "NWM": "Non Woven Print Module", "NWPR": "Nw Printed Module",
        "NWU": "Non Wowen Unwinder", "OCU": "Outer Cover Unwind",
        "OLP": "On Line Printed", "PCU": "Pouch Unwind", "PPU": "Pulp Unwind",
        "RPU": "Release Paper Unwind", "SLU": "Surge Layer Unwind",
        "SPU": "Side Panel Unwind", "STU": "Spun Tissue Unwind",
        "SWU": "Solo Wrap Unwind", "TLU": "Transfer Layer Unwind",
        "TSU": "Topsheet Unwind", "TTU": "Debobinador De Tissue",
        "WEU": "Waist Elastic Unwind", "WLU": "Wrap Layer Unwind",
        "WTU": "Wings Tape Unwind",
    },
    "Empaque / Paletizado / Transporte": {
        "BG1": "Bagger1", "BG2": "Bagger 2", "BGR": "Bagger 2", "CNV": "Conveyors",
        "CSP": "Casepacker", "CTN": "Cartonner", "DBD": "D-Bund", "GRC": "Grecon",
        "MPK": "Multipacker", "PHS": "Sistema De Empaque Manual",
        "SPK": "Empaque Secundario", "STK": "Stacker",
    },
    "Otros / sin mapear": {
        "ETA": "Aplicador De Etiqueta", "MEA": "Main Drive",
        "SEA": "Aplicador Oreja Posterior", "TRA": "Zona Transmision",
    },
    "Sistemas centrales / eléctricos": {
        "ACS": "Sistema De Ventiladores", "AUX": "Auxiliary System",
        "CCM": "Machine Fans", "CS": "Central System", "ELE": "Electrico",
        "FRT": "Filtro Rotativo", "LIBRE": "Sin Clasificar",
        "MACT": "Actividades De Maquina", "MCC": "Machine Center Control",
        "MIT": "Machine Interfase Technical", "MNPL": "Main Electrical Cabinets",
        "OP": "Machine Operator Panel", "OTRO": "Conveyor S",
        "PP1": "Processor Panel", "RIS": "Tablero Ris", "TBL": "Tableros Electricos",
        "TSV": "Tablero Servos",
    },
}

# --- Reglas por palabras clave (texto en minúsculas y sin tildes).
#     Se prueban en orden; gana la primera que coincide.
TIPO_REGLAS = [
    ("Sin asignar",         r"unassigned|sin asignar|no asignad"),
    ("Eléctricas",          r"electr|\belec\b"),
    ("Mecánicas",           r"mecan|mechan|\bmec\b"),
    ("TNP",                 r"\btnp\b|no programad|schedule loss"),
    ("Paradas Programadas", r"programad|planificad|scheduled|planned"),
    ("Cambio de Grado",     r"cambio|grado|grade|changeover"),
    ("Causas Externas",     r"extern|^operations$"),
    ("Utilities",           r"utilit"),
    ("Paradas Menores",     r"menor|minor|operat|process|proceso|libre"),
]

FAMILIA_REGLAS = [
    ("No clasificado / Libre",
     r"^(libre|\(sin dato\)|otros?|sin dato)$|atoro|material ausente|"
     r"ausencia de material|descarte|desecho"),
    ("Seguridad (guardas / pilz / paros)",
     r"pilz|safety|seguridad|guarda|seguro de puerta|boton de stop|emergencia|"
     r"botonera"),
    ("Comunicaciones / Red",
     r"comunicac|communicat|controlnet|ethernet|\bred\b|\bnodo\b|network|"
     r"profibus|devicenet"),
    ("Sistema de Control (PLC/lógica)",
     r"control fault|falla de control|sistema de control|\bplc\b|\bhmi\b|"
     r"power fault|procesador|\bcmp\d|\bmmc\b|\bi/o\b"),
    ("Guiado y control de Web (FIFE)",
     r"fife|control de web|guiado de web|^web$|web guide"),
    ("Empaque / Paletizado",
     r"bagger (?!conveyor)|bolsa|pusher|empujador|robot|paletiz|colero|flipper|"
     r"caja|magazine|abridor|barra de|brazo de succion|empaque(?! individual)"),
    ("Servos / Drives / Motores",
     r"servo|drives?\b|driver|motores|motors"),
    ("Sensores / Visión / Cámaras",
     r"sensor|vision|camara (checker|sli|fod|vision)|checker|print head|balanza|"
     r"detector|encoder|inspeccion"),
    ("Sistema Neumático / Aire",
     r"neumatic|pneumat|presion de aire|vacuum|motor vacio|actuador|valvula|"
     r"valve|manifold|aire comprimido"),
    ("Fajas / Bandas / Transporte",
     r"faja|banda|belt|conveyor|transportador|transporte|cadena|chain|infeed|"
     r"traslador"),
    ("Aplicadores / Manejo de material",
     r"aplicador|applicator|debobinador|desenrrollador|empalmador|festoon|"
     r"acumulador|rodillo|nip roll|tambor|bonder|anvil|[bv]alerino|casetera"),
    ("Servicios / Auxiliares",
     r"ventilador|\bfan\b|melter|manguera|hoses|chiller|generat|fiberizer|"
     r"molino|filtro rotativo|calentamiento"),
    ("Corte / Sellado / Compresión",
     r"corte|\bcut|knife|cuchill|seal|sellad|soldad|termosell|compres|"
     r"press (pad|unit)|embossing unit|\(dado\)|blade|pressure switch"),
    ("Eléctrico / Electrónico (genérico)",
     r"^electr|caida de tension|pico de tension|sistema electrico|voltaje"),
]
FAMILIA_OTROS = "Otros componentes"

AREA_REGLAS = [
    ("Sistemas centrales / eléctricos",
     r"^libre\b|sin clasificar|central|electric|tablero|cabinet|(?<!side )panel|"
     r"control|ventilador|\bfans?\b|auxiliar|operator|processor|interfase|"
     r"interface|filtro rotativo|actividades"),
    ("Adhesivos / Melters", r"adhesiv|melter|glue|hot ?melt"),
    ("Debobinadores / Unwinds",
     r"unwind|debobinador|desenrr?ollador|printed|print module|on line print"),
    ("Empaque / Paletizado / Transporte",
     r"bagger|packer|pack|empaque|carton|stacker|conveyor|palet|d-?bund|grecon"),
    ("Aplicadores / Formación de producto", r"pouch|aplicador|applicat"),
    ("Corte y sellado",
     r"cut|corte|seal|sell|press|prensa|trifold|coupler|repitch|rotator|"
     r"wrap module"),
    ("Aplicadores / Formación de producto",
     r"application|forming|formador|formacion|folding|embossing|embosser|"
     r"debulker|fiberizer|mill|absorbent|airlaid|perfume|rodillo|ultrasonic|"
     r"elastic|borning|wrap"),
]
AREA_OTROS = "Otros / sin mapear"


# =============================================================================
# 2. UTILIDADES DE TEXTO
# =============================================================================
VACIOS = {"", "nan", "none", "null", "nat", "-", "--", "n/a", "#n/a", "na",
          "(sin dato)", "sin dato"}


def es_nulo(v):
    if v is None:
        return True
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def clave(s):
    """Clave para comparar: minúsculas, sin tildes, solo letras y números."""
    if es_nulo(s):
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s.lower())


def plano(s):
    """Minúsculas y sin tildes, conservando espacios: es lo que leen las reglas."""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = re.sub(r"[^a-z0-9#/&().,+\-]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def limpiar(v):
    """Texto sin espacios de más; None si es un vacío ("", "nan", "-", ...)."""
    if es_nulo(v):
        return None
    s = unicodedata.normalize("NFC", str(v))
    s = s.replace("\u00a0", " ").replace("\t", " ").replace("\r", " ").replace("\n", " ")
    s = re.sub(r"\s+", " ", s).strip().strip("\"'").strip()
    return None if s.lower() in VACIOS else s


def capitalizar(s):
    return s[:1].upper() + s[1:] if s and s[:1].islower() else s


def titulo(s):
    """'ADHESIVO DE CONSTRUCCION' -> 'Adhesivo De Construccion' (como el histórico)."""
    return re.sub(r"[^\W\d_]+", lambda m: m.group(0)[:1].upper() + m.group(0)[1:].lower(),
                  str(s))


def por_reglas(texto, reglas, defecto):
    t = plano(texto)
    for destino, patron in reglas:
        if re.search(patron, t):
            return destino
    return defecto


def moda(serie):
    s = serie.dropna()
    return s.mode().iloc[0] if len(s) else None


def por_valor(serie, func):
    """Aplica func una sola vez por valor distinto (rápido con 100 mil filas)."""
    codigos, unicos = pd.factorize(serie)
    vals = np.empty(len(unicos) + 1, dtype=object)
    vals[:-1] = [func(u) for u in unicos]
    vals[-1] = func(None)                       # código -1 = vacío
    return pd.Series(vals[codigos], index=serie.index, dtype=object)


def mas_frecuente(claves, valores):
    """{clave: valor más frecuente} sin agrupar fila por fila."""
    t = pd.DataFrame({"k": np.asarray(claves, dtype=object),
                      "v": np.asarray(valores, dtype=object)}).dropna()
    t = t[t["k"] != ""]
    t = t.groupby(["k", "v"]).size().reset_index(name="n")
    t = t.sort_values("n", ascending=False, kind="stable").drop_duplicates("k")
    return dict(zip(t["k"], t["v"]))


def _ruta(p):
    """Busca la ruta tal cual y, si no está, junto a este script."""
    if not p or os.path.isabs(p) or os.path.exists(p):
        return p
    junto = os.path.join(os.path.dirname(os.path.abspath(__file__)), p)
    return junto if os.path.exists(junto) else p


# =============================================================================
# 3. LECTURA DEL EXPORT  (tolerante a cambios de columnas)
# =============================================================================
ALIAS_ENTRADA = {
    "unidad":  ["pudesc", "unidad", "unit", "pu", "equipo"],
    "inicio":  ["fechainicioreal", "fechainicio", "inicio", "starttime", "start",
                "fechahora", "fecha", "timestamp"],
    "fin":     ["fechafinreal", "fechafin", "fin", "endtime", "end"],
    "turno":   ["turno", "shift"],
    "n1":      ["reasonlevel1", "nivel1", "razon1", "level1", "reason1"],
    "n2":      ["reasonlevel2", "nivel2", "razon2", "level2", "reason2",
                "reasonlevel12", "modulo"],
    "n3":      ["reasonlevel3", "nivel3", "razon3", "level3", "reason3",
                "sistemal3"],
    "n4":      ["reasonlevel4", "nivel4", "razon4", "level4", "reason4",
                "detallel4"],
    "delay":   ["delay", "horas", "duracion", "duration", "hrs", "downtime"],
    "maquina": ["maquina", "machine", "maq"],
    "tipo":    ["tipo", "type", "categoria", "category"],
}


def mapear_entrada(columnas):
    norm = {}
    for c in columnas:
        norm.setdefault(clave(c), c)
    usadas, out = set(), {}
    for destino, opciones in ALIAS_ENTRADA.items():
        hallado = next((norm[o] for o in opciones
                        if o in norm and norm[o] not in usadas), None)
        if hallado is None:                       # coincidencia parcial
            hallado = next((orig for k, orig in norm.items() if orig not in usadas
                            and any(len(o) >= 5 and k.startswith(o) for o in opciones)),
                           None)
        if hallado is not None:
            out[destino] = hallado
            usadas.add(hallado)
    return out


def leer_tabla(ruta, hoja=None):
    if ruta.lower().endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "latin-1"):
            try:
                return pd.read_csv(ruta, sep=None, engine="python", encoding=enc)
            except UnicodeDecodeError:
                continue
    xl = None
    if importlib.util.find_spec("python_calamine"):   # "pip install python-calamine"
        try:                                           # lee Excel mucho más rápido
            xl = pd.ExcelFile(ruta, engine="calamine")
        except (ImportError, ValueError):
            xl = None
    xl = xl or pd.ExcelFile(ruta)
    hoja = hoja if hoja in xl.sheet_names else xl.sheet_names[0]
    return xl.parse(hoja)


def _a_datetime(serie, dayfirst):
    try:
        return pd.to_datetime(serie, errors="coerce", dayfirst=dayfirst, format="mixed")
    except (TypeError, ValueError):                     # pandas < 2.0
        return pd.to_datetime(serie, errors="coerce", dayfirst=dayfirst)


def a_fecha(serie):
    """Fecha de Excel, número de serie, texto dd/mm/aaaa o ISO aaaa-mm-dd."""
    if pd.api.types.is_datetime64_any_dtype(serie):
        return serie
    if pd.api.types.is_numeric_dtype(serie):            # número de serie de Excel
        return pd.to_datetime(serie, unit="D", origin="1899-12-30", errors="coerce")
    txt = serie.map(lambda v: v if isinstance(v, (datetime, pd.Timestamp))
                    else limpiar(v))
    # con dayfirst=True pandas lee "2026-01-06" como 1 de junio: el ISO va aparte
    iso = txt.map(lambda v: isinstance(v, str)
                  and bool(re.match(r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", v)))
    fecha = _a_datetime(txt.where(~iso), dayfirst=True)
    if iso.any():
        fecha = fecha.where(~iso, _a_datetime(txt.where(iso), dayfirst=False))
    return fecha


# =============================================================================
# 4. HISTORICO Y CATALOGOS
# =============================================================================
class Catalogo:
    """Módulos (código -> nombre, área) y familias (sistema L3 -> familia)."""

    def __init__(self):
        self.mod = {}          # COD -> (nombre, area, origen)
        for area, mods in MODULOS_BASE.items():
            for cod, nom in mods.items():
                self.mod[cod.upper()] = (nom, area, "base")
        self.fam = {}          # clave(L3) -> familia (histórico)
        self.eventos = {}      # clave de evento -> COD
        self.eventos2 = {}     # (unidad, fecha, tipo) -> COD si es único
        self.fuentes = []

    def aplicar_manuales(self):
        for cod, val in MODULOS_MANUAL.items():
            val = [val] if isinstance(val, str) else list(val)   # "COD": "Nombre"
            nom, area = (val + [None, None])[:2]
            ant = self.mod.get(cod.upper(), (cod.upper(), AREA_OTROS, ""))
            self.mod[cod.upper()] = (nom or ant[0], area or ant[1], "manual")
        self._nombres()

    def _nombres(self):
        self.por_nombre = {}
        for cod, (nom, _, _) in self.mod.items():
            self.por_nombre.setdefault(clave(nom), cod)
            self.por_nombre.setdefault(clave(f"{cod} {nom}"), cod)
        for txt, cod in ALIAS_MODULOS.items():
            self.por_nombre[clave(txt)] = cod.upper()
        self.claves_nombre = [k for k in self.por_nombre if len(k) >= 6]

    def nombre_area(self, cod):
        nom, area, _ = self.mod[cod]
        return nom, area


SEP_MODULO = r"\s+[—–-]\s+"          # "AS1 — Adhesivo De Construccion"


def _col(df, *opciones):
    norm = {clave(c): c for c in df.columns}
    return next((norm[o] for o in opciones if o in norm), None)


def clave_evento(unidad, fecha, l3, l4, tipo):
    """k4 = unidad|inicio|L3|L4 (exacta).  k2 = unidad|inicio|tipo (respaldo)."""
    f = pd.to_datetime(fecha, errors="coerce").dt.round("s").dt.strftime("%Y-%m-%d %H:%M:%S")
    base = por_valor(unidad, clave) + "|" + f.fillna("")
    k4 = (base + "|" + por_valor(l3, lambda v: clave(v) or "sindato") + "|" +
          por_valor(l4, lambda v: clave(v) or "sindato"))
    k2 = base + "|" + por_valor(tipo, lambda v: clave(estandarizar_tipo(v)))
    return k4, k2


def cargar_historico(cat, avisos):
    for ruta, hoja in HISTORICO:
        ruta = _ruta(ruta)
        if not ruta or not os.path.exists(ruta):
            avisos.append(f"No encontré el histórico '{ruta}'. Se omite.")
            continue
        try:
            h = leer_tabla(ruta, hoja)
        except Exception as e:                       # archivo abierto, dañado...
            avisos.append(f"No pude leer el histórico '{ruta}' ({e}).")
            continue
        c = dict(unidad=_col(h, "unidad"), fecha=_col(h, "fecha"),
                 cod=_col(h, "cod"), modulo=_col(h, "modulo"), area=_col(h, "area"),
                 familia=_col(h, "familia"),
                 l3=_col(h, "l3original", "sistemal3", "sistema", "l3"),
                 l4=_col(h, "l4original", "detallel4", "detalle", "l4"),
                 l3std=_col(h, "sistemal3"), tipo=_col(h, "tipo"))
        if not (c["cod"] or c["modulo"]):
            avisos.append(f"El histórico '{ruta}' no tiene columnas Cod/Modulo. Se omite.")
            continue
        h = h.copy()
        partes = (por_valor(h[c["modulo"]], limpiar).str.split(SEP_MODULO, n=1, regex=True)
                  if c["modulo"] else None)
        cod = h[c["cod"]] if c["cod"] else partes.str[0]
        h["_cod"] = por_valor(cod, limpiar).str.upper()
        h.loc[h["_cod"].isin([SIN.upper(), "SINDATO"]), "_cod"] = None
        h["_nom"] = partes.str[1] if c["modulo"] else None
        h["_area"] = por_valor(h[c["area"]], limpiar) if c["area"] else None

        # módulos: código -> nombre y área más frecuentes
        n_mod = 0
        for cod_, g in h.dropna(subset=["_cod"]).groupby("_cod"):
            nom = moda(g["_nom"]) if c["modulo"] else None
            area = moda(g["_area"]) if c["area"] else None
            ant = cat.mod.get(cod_)
            if ant and ant[2] == "manual":
                continue
            cat.mod[cod_] = (nom or (ant[0] if ant else cod_),
                             area or (ant[1] if ant else AREA_OTROS), "histórico")
            n_mod += 1

        # familias: sistema L3 -> familia más frecuente
        if c["familia"] and (c["l3"] or c["l3std"]):
            for col in {c["l3"], c["l3std"]} - {None}:
                fam = mas_frecuente(por_valor(h[col], clave),
                                    por_valor(h[c["familia"]], limpiar))
                for k, f in fam.items():
                    if k and f and k not in cat.fam:
                        cat.fam[k] = f

        # eventos: (unidad, fecha, L3, L4) -> código de módulo
        n_ev = 0
        if c["unidad"] and c["fecha"]:
            l3 = h[c["l3"]] if c["l3"] else pd.Series(None, index=h.index)
            l4 = h[c["l4"]] if c["l4"] else pd.Series(None, index=h.index)
            tipo = h[c["tipo"]] if c["tipo"] else pd.Series(None, index=h.index)
            k4, k2 = clave_evento(h[c["unidad"]], h[c["fecha"]], l3, l4, tipo)
            ev = pd.DataFrame({"k4": k4, "k2": k2, "cod": h["_cod"]}).dropna(subset=["cod"])
            for k, cod_ in mas_frecuente(ev["k4"], ev["cod"]).items():
                cat.eventos.setdefault(k, cod_)
            if c["tipo"]:          # respaldo si L3/L4 cambiaron: misma unidad, hora y tipo
                distintos = ev.groupby("k2")["cod"].nunique()
                unicos = (ev[ev["k2"].isin(distintos[distintos == 1].index)]
                          .drop_duplicates("k2"))
                for k, cod_ in zip(unicos["k2"], unicos["cod"]):
                    cat.eventos2.setdefault(k, cod_)
            n_ev = len(ev)
        cat.fuentes.append(f"{os.path.basename(ruta)} [{hoja}]: {len(h):,} filas, "
                           f"{n_mod} módulos, {n_ev:,} eventos con módulo")


# =============================================================================
# 5. ESTANDARIZACION
# =============================================================================
def maquina_valida(v):
    s = limpiar(v)
    if not s:
        return None
    m = re.fullmatch(r"([A-Z]{1,4})-?(\d{1,3})", re.sub(r"[\s_]", "", s.upper()))
    return m.group(1) + m.group(2) if m else None


def maquina_de_unidad(u):
    m = re.match(r"^\s*([A-Za-z]{1,4})\s*-?\s*(\d{1,3})(?!\d)", str(u or ""))
    return (m.group(1) + m.group(2)).upper() if m else None


def estandarizar_tipo(v):
    s = limpiar(v)
    return por_reglas(s, TIPO_REGLAS, None) if s else None


def _quitar_prefijo(s):
    if not s or not PREFIJOS_L4:
        return s
    pat = r"^\s*(?:%s)\s*-\s*(?=\S)" % "|".join(re.escape(p) for p in PREFIJOS_L4)
    return re.sub(pat, "", s, flags=re.I)


def estandarizar_texto(serie, quitar_prefijo=False):
    """Devuelve la serie estándar y el mapa {original limpio: estándar}."""
    manual = {clave(k): v for k, v in TEXTOS_MANUAL.items()}
    cuenta = por_valor(serie, limpiar).value_counts()
    paso = {}                                    # original limpio -> intermedio
    peso = Counter()
    for orig, n in cuenta.items():
        t = _quitar_prefijo(orig) if quitar_prefijo else orig
        t = manual.get(clave(orig)) or manual.get(clave(t)) or t
        paso[orig] = t
        peso[t] += n
    mejor = {}
    if UNIFICAR_VARIANTES:                       # la escritura más frecuente gana
        orden = sorted(peso.items(), key=lambda kv: (-kv[1], -sum(ord(ch) > 127
                                                                  for ch in kv[0]), kv[0]))
        for t, _ in orden:
            mejor.setdefault(clave(t), t)
    final = {o: capitalizar(mejor.get(clave(t), t)) for o, t in paso.items()}
    salida = por_valor(serie, lambda v: final.get(limpiar(v), SIN) if limpiar(v) else SIN)
    return salida, final


def parsear_modulo(txt, maquina, cat):
    """Texto del nivel 2 -> (código, nombre leído). Código None si no se reconoce."""
    t = limpiar(txt)
    if not t:
        return None, None
    if maquina and maquina != SIN:                   # "PI8 - AS1 - ..." -> "AS1 - ..."
        resto = re.sub(rf"^{re.escape(maquina)}(?![A-Za-z0-9])\s*[-–—:|_]*\s*", "", t,
                       flags=re.I).strip()
        t = resto or t
    k = clave(t)
    if k in cat.por_nombre:                          # nombre o "COD nombre" conocido
        return cat.por_nombre[k], None
    if k in ("libre", "sinclasificar", "noclasificado", "sinclasificacion"):
        return "LIBRE", None

    def es_cod(c):
        return c.upper() in cat.mod or c.isupper() or bool(re.search(r"\d", c))

    pats = [r"^\(?([A-Za-z][A-Za-z0-9]{1,5})\)?\s*[-–—:|]+\s*(.+)$",   # AS1 - Nombre
            r"^\(([A-Za-z][A-Za-z0-9]{1,5})\)\s*(.+)$"]                # (AS1) Nombre
    for p in pats:
        m = re.match(p, t)
        if m and es_cod(m.group(1)):
            return m.group(1).upper(), m.group(2).strip()
    m = re.match(r"^(.+?)\s*\(([A-Za-z][A-Za-z0-9]{1,5})\)$", t)     # Nombre (AS1)
    if m and (m.group(2).upper() in cat.mod or m.group(2).isupper()):
        return m.group(2).upper(), m.group(1).strip()
    m = re.match(r"^([A-Za-z][A-Za-z0-9]{1,5})\s+(.+)$", t)           # AS1 Nombre
    if m and (m.group(1).upper() in cat.mod or
              (m.group(1).isupper() and re.search(r"\d", m.group(1)))):
        return m.group(1).upper(), m.group(2).strip()
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9]{1,5}", t) and (t.upper() in cat.mod
                                                         or t.isupper()):
        return t.upper(), None                                         # solo código
    cerca = difflib.get_close_matches(k, cat.claves_nombre, n=1, cutoff=0.9)
    if cerca:
        return cat.por_nombre[cerca[0]], None
    return None, t


def nivel2_util(r):
    """¿La columna de nivel 2 trae módulos? (y no es copia del nivel 1)."""
    if "n2" not in r:
        return False, "el export no tiene columna de nivel 2 (módulo)"
    b = por_valor(r["n2"], clave)
    llenos = b != ""
    if llenos.mean() < 0.01:
        return False, f"la columna '{r.attrs['cols']['n2']}' viene vacía"
    if "n1" in r:
        iguales = (por_valor(r["n1"], clave) == b) & llenos
        if iguales.sum() > 0.95 * llenos.sum():
            return False, (f"la columna '{r.attrs['cols']['n2']}' es una copia exacta de "
                           f"'{r.attrs['cols']['n1']}', no trae el módulo")
    return True, ""


def asignar_modulos(d, cat, usar_n2, log):
    cod = pd.Series(None, index=d.index, dtype=object)
    leido = pd.Series(None, index=d.index, dtype=object)
    origen = pd.Series("Sin dato", index=d.index, dtype=object)

    if usar_n2:
        n2 = por_valor(d["n2"], limpiar)                # None si viene vacío (no NaN)
        res = {p: parsear_modulo(p[0], p[1], cat) for p in set(zip(n2, d["Maquina"]))}
        par = [res[p] for p in zip(n2, d["Maquina"])]
        cod[:] = [p[0] for p in par]
        leido[:] = [p[1] for p in par]
        origen[cod.notna() | leido.notna()] = "Nivel 2"

    falta = cod.isna() & leido.isna()
    if falta.any() and (cat.eventos or cat.eventos2):
        k4, k2 = clave_evento(d.loc[falta, "Unidad"], d.loc[falta, "Fecha"],
                              d.loc[falta, "n3"], d.loc[falta, "n4"], d.loc[falta, "Tipo"])
        rec = k4.map(cat.eventos)
        rec = rec.where(rec.notna(), k2.map(cat.eventos2))
        cod[falta] = rec
        origen[falta & cod.notna()] = "Histórico"

    nuevos = {}
    cod_f, mod_f, area_f = [], [], []
    for c, nombre_leido in zip(cod, leido):
        c = None if es_nulo(c) else c                 # NaN cuenta como verdadero
        nombre_leido = None if es_nulo(nombre_leido) else nombre_leido
        if c and c in cat.mod:
            nom, area = cat.nombre_area(c)
        elif c or nombre_leido:                       # código o texto no catalogado
            c = c or re.sub(r"\s+", " ", plano(nombre_leido).upper())
            nom = titulo(nombre_leido) if nombre_leido else c
            area = por_reglas(f"{c} {nom}", AREA_REGLAS, AREA_OTROS)
            nuevos.setdefault(c, (nom, area))
        else:
            cod_f.append(SIN), mod_f.append(SIN), area_f.append(SIN)
            continue
        cod_f.append(c), mod_f.append(f"{c} — {nom}"), area_f.append(area)
    d["Cod"], d["Modulo"], d["Area"], d["Origen módulo"] = cod_f, mod_f, area_f, origen
    log["mod_nuevos"] = nuevos
    return d


def asignar_familias(d, cat):
    manual = {clave(k): v for k, v in FAMILIAS_MANUAL.items()}
    resultado, origen = {}, {}
    for std, orig in d[["Sistema (L3)", "L3 original"]].drop_duplicates().itertuples(
            index=False):
        kk = [clave(std), clave(orig)]
        regla = (por_reglas(std, FAMILIA_REGLAS, FAMILIA_OTROS) if std != SIN
                 else "No clasificado / Libre")
        fam = next((manual[k] for k in kk if k in manual), None)
        if fam:
            o = "Manual"
        else:
            hist = next((cat.fam[k] for k in kk if k in cat.fam), None)
            if hist and not (RECLASIFICAR_OTROS and hist == FAMILIA_OTROS
                             and regla != FAMILIA_OTROS):
                fam, o = hist, "Histórico"
            elif hist:
                fam, o = regla, f"Regla (antes {FAMILIA_OTROS})"
            else:
                fam, o = regla, "Regla"
        resultado[(std, orig)], origen[(std, orig)] = fam, o
    llaves = list(zip(d["Sistema (L3)"], d["L3 original"]))
    d["Familia"] = [resultado[k] for k in llaves]
    d["_origen_familia"] = [origen[k] for k in llaves]
    return d


def estandarizar(ruta, cat, avisos):
    log = {}
    crudo = leer_tabla(ruta, HOJA_ENTRADA)
    crudo.columns = [str(c).strip() for c in crudo.columns]
    crudo = crudo.dropna(how="all")
    cols = mapear_entrada(crudo.columns)
    log["columnas"] = cols
    falta = [k for k in ("inicio",) if k not in cols]
    if "unidad" not in cols and "maquina" not in cols:
        falta.append("unidad o maquina")
    if "delay" not in cols and "fin" not in cols:
        falta.append("delay o fecha fin")
    if falta:
        raise SystemExit(f"\n  No encuentro columnas obligatorias: {falta}\n"
                         f"  Columnas del archivo: {list(crudo.columns)}\n")

    r = pd.DataFrame({k: crudo[c] for k, c in cols.items()})
    r.attrs["cols"] = cols
    log["leidas"] = len(r)

    # --- duplicados exactos
    log["duplicados"] = 0
    if QUITAR_DUPLICADOS:
        dup = r.duplicated()
        log["duplicados"] = int(dup.sum())
        r = r[~dup].copy()
        r.attrs["cols"] = cols

    usar_n2, motivo = nivel2_util(r)
    log["usa_n2"], log["motivo_n2"] = usar_n2, motivo

    d = pd.DataFrame(index=r.index)
    obs = pd.Series("", index=r.index, dtype=object)

    def anotar(mask, texto):
        obs.loc[mask] = obs.loc[mask].map(lambda s: f"{s}; {texto}" if s else texto)

    # --- unidad y máquina
    d["Unidad"] = (por_valor(r["unidad"], limpiar) if "unidad" in r
                   else pd.Series(None, index=r.index, dtype=object))
    maq = (por_valor(r["maquina"], maquina_valida) if "maquina" in r
           else pd.Series(None, index=r.index, dtype=object))
    de_unidad = por_valor(d["Unidad"], maquina_de_unidad)
    rescate = maq.isna() & de_unidad.notna()
    log["maq_rescatada"] = int(rescate.sum())
    anotar(rescate, "máquina tomada de la unidad")
    d["Maquina"] = maq.where(maq.notna(), de_unidad)
    log["maq_sin"] = int(d["Maquina"].isna().sum())
    log["maq_sin_unidades"] = ", ".join(
        f"{u} ({n:,})" for u, n in d.loc[d["Maquina"].isna(), "Unidad"].value_counts()
        .head(5).items())
    d["Maquina"] = d["Maquina"].fillna(SIN)
    d["Unidad"] = d["Unidad"].where(d["Unidad"].notna(), d["Maquina"]).fillna(SIN)

    # --- fechas y duración
    d["Fecha"] = a_fecha(r["inicio"])
    d["Fin"] = a_fecha(r["fin"]) if "fin" in r else pd.NaT
    dur = (d["Fin"] - d["Fecha"]).dt.total_seconds() / 3600
    dur = dur.where(dur >= 0)
    delay = (pd.to_numeric(r["delay"], errors="coerce") if "delay" in r
             else pd.Series(np.nan, index=r.index))
    ambos = delay.notna() & dur.notna() & (dur > 0) & (delay > 0)
    log["delay_unidad"] = "horas"
    if ambos.sum() >= 20:                          # ¿Delay en horas, minutos o segundos?
        ratio = float((delay[ambos] / dur[ambos]).median())
        if 40 < ratio < 80:
            delay, log["delay_unidad"] = delay / 60, "minutos (convertido a horas)"
        elif 2400 < ratio < 4800:
            delay, log["delay_unidad"] = delay / 3600, "segundos (convertido a horas)"
    calc = delay.isna() & dur.notna()
    log["dur_calculada"] = int(calc.sum())
    anotar(calc, "horas calculadas con Fin - Inicio")
    difiere = ambos & ((delay - dur).abs() > 1 / 60)
    log["delay_difiere"] = int(difiere.sum())
    anotar(difiere, "Delay distinto de Fin - Inicio")
    horas = delay.where(delay.notna(), dur)
    log["sin_duracion"] = int(horas.isna().sum())
    anotar(horas.isna(), "sin duración")
    log["negativas"] = int((horas < 0).sum())
    d["Horas"] = horas.fillna(0).clip(lower=0)

    # --- turno
    d["Turno"] = (por_valor(r["turno"], lambda v: re.sub(r"^turno\s*", "", limpiar(v),
                                                          flags=re.I).upper()
                            if limpiar(v) else SIN)
                  if "turno" in r else SIN)

    # --- tipo
    tipo = (por_valor(r["tipo"], estandarizar_tipo) if "tipo" in r
            else pd.Series(None, index=r.index, dtype=object))
    desde_n1 = tipo.isna() | (tipo == "Sin asignar")
    if "n1" in r:
        alt = por_valor(r["n1"], estandarizar_tipo)
        usar = desde_n1 & alt.notna() & (alt != "Sin asignar")
        tipo[usar] = alt[usar]
        anotar(usar, "tipo tomado del nivel 1")
    sin_regla = tipo.isna() & ((por_valor(r["tipo"], limpiar).notna() if "tipo" in r
                                else False) |
                               (por_valor(r["n1"], limpiar).notna() if "n1" in r
                                else False))
    texto_tipo = (r["tipo"] if "tipo" in r else r.get("n1", pd.Series(None, index=r.index)))
    log["tipo_sin_regla"] = (texto_tipo[sin_regla].map(limpiar).value_counts()
                             if sin_regla.any() else pd.Series(dtype=int))
    tipo[sin_regla] = texto_tipo[sin_regla].map(limpiar)
    d["Tipo"] = tipo.fillna("Sin asignar")
    d["N1 original"] = por_valor(r["n1"], limpiar).fillna(SIN) if "n1" in r else SIN

    # --- textos L3 / L4
    vacia = pd.Series(None, index=r.index, dtype=object)
    d["Sistema (L3)"], log["map_l3"] = estandarizar_texto(r.get("n3", vacia))
    d["Detalle (L4)"], log["map_l4"] = estandarizar_texto(r.get("n4", vacia),
                                                          quitar_prefijo=QUITAR_PREFIJO_L4)
    d["L3 original"] = por_valor(r.get("n3", vacia), limpiar).fillna(SIN)
    d["L4 original"] = por_valor(r.get("n4", vacia), limpiar).fillna(SIN)
    d["n2"], d["n3"], d["n4"] = r.get("n2", vacia), r.get("n3", vacia), r.get("n4", vacia)
    if usar_n2:
        d["N2 original"] = por_valor(r["n2"], limpiar).fillna(SIN)
    d["Observaciones"] = obs

    # --- filtros
    log["sin_fecha"] = int(d["Fecha"].isna().sum())
    d = d[d["Fecha"].notna()]
    log["por_tipo"] = d["Tipo"].value_counts()
    if TIPOS_INCLUIR:
        incluir = {clave(t) for t in TIPOS_INCLUIR}
        d = d[por_valor(d["Tipo"], clave).isin(incluir)]
    log["tras_tipo"] = len(d)
    if FECHA_DESDE:
        d = d[d["Fecha"] >= pd.Timestamp(FECHA_DESDE)]
    if FECHA_HASTA:
        d = d[d["Fecha"] < pd.Timestamp(FECHA_HASTA) + pd.Timedelta(days=1)]
    log["fuera_fecha"] = log["tras_tipo"] - len(d)
    d = d.copy()

    # --- módulo, área y familia
    d = asignar_modulos(d, cat, usar_n2, log)
    d = asignar_familias(d, cat)

    # --- columnas finales
    d["Mes"] = d["Fecha"].dt.month
    d["Año"] = d["Fecha"].dt.year
    d["Minutos"] = d["Horas"] * 60
    d = d.sort_values(["Fecha", "Maquina", "Unidad"], kind="stable").reset_index(drop=True)
    return d, log


# =============================================================================
# 6. HOJAS DE CONTROL
# =============================================================================
def hoja_calidad(d, log, cat, avisos):
    filas = [("AVISO", "", a) for a in avisos]
    c = log["columnas"]
    filas += [
        ("Columnas reconocidas", len(c), ", ".join(f"{k}={v}" for k, v in c.items())),
        ("Filas leídas", log["leidas"], ""),
        ("Duplicados exactos eliminados", log["duplicados"], ""),
        ("Delay venía en", "", log["delay_unidad"]),
        ("Horas calculadas con Fin - Inicio", log["dur_calculada"],
         "Delay vacío en el export"),
        ("Delay distinto de Fin - Inicio (>1 min)", log["delay_difiere"],
         "se respeta el Delay"),
        ("Filas sin duración (quedan con 0 h)", log["sin_duracion"], ""),
        ("Máquina tomada de la unidad", log["maq_rescatada"],
         "columna Maquina vacía o inválida (ej. 'Even')"),
        ("Máquina sin identificar", log["maq_sin"], log["maq_sin_unidades"]),
        ("Filas sin fecha de inicio (descartadas)", log["sin_fecha"], ""),
    ]
    for t, n in log["por_tipo"].items():
        dentro = not TIPOS_INCLUIR or clave(t) in {clave(x) for x in TIPOS_INCLUIR}
        filas.append(("Tipo: " + str(t), int(n), "se incluye" if dentro else "excluido"))
    filas += [
        ("Fuera del rango de fechas", log["fuera_fecha"],
         f"desde {FECHA_DESDE or 'inicio'} hasta {FECHA_HASTA or 'fin'}"),
        ("FILAS FINALES", len(d), f"{d['Horas'].sum():,.1f} h"),
        ("Módulo: nivel 2 del export", int((d["Origen módulo"] == "Nivel 2").sum()),
         "" if log["usa_n2"] else log["motivo_n2"]),
        ("Módulo: recuperado del histórico",
         int((d["Origen módulo"] == "Histórico").sum()), "; ".join(cat.fuentes)),
        ("Módulo: sin dato", int((d["Origen módulo"] == "Sin dato").sum()),
         "ni el export ni el histórico lo traen"),
    ]
    for o, n in d["_origen_familia"].value_counts().items():
        filas.append((f"Familia: {o}", int(n), ""))
    for campo, mapa, col in (("L3", log["map_l3"], "Sistema (L3)"),
                             ("L4", log["map_l4"], "Detalle (L4)")):
        cambiados = sum(1 for o, s in mapa.items() if o != s)
        filas.append((f"{campo}: textos distintos", d[col].nunique(),
                      f"{len(mapa)} escrituras originales, {cambiados} corregidas "
                      f"(ver hoja Unificaciones)"))
    return pd.DataFrame(filas, columns=["Paso", "Cantidad", "Detalle"])


def hoja_pendientes(d, log, cat):
    filas = []

    def agrega(tema, mask, por, asignado, nota):
        g = (d[mask].groupby(por).agg(Eventos=("Horas", "size"), Horas=("Horas", "sum"))
             .reset_index().sort_values("Horas", ascending=False))
        for _, x in g.iterrows():
            filas.append((tema, x[por], asignado(x), int(x["Eventos"]),
                          round(float(x["Horas"]), 2), nota))

    of = d["_origen_familia"]
    fam = d.drop_duplicates("Sistema (L3)").set_index("Sistema (L3)")["Familia"]
    agrega("Familia asignada por regla", of == "Regla", "Sistema (L3)",
           lambda x: fam[x["Sistema (L3)"]],
           "sistema nuevo: confirma la familia o corrígela en FAMILIAS_MANUAL")
    agrega("Familia reclasificada", of.str.startswith("Regla (antes"), "Sistema (L3)",
           lambda x: fam[x["Sistema (L3)"]],
           f"tu histórico lo tenía en '{FAMILIA_OTROS}'. RECLASIFICAR_OTROS=False lo deja")
    nuevos = log.get("mod_nuevos", {})
    if nuevos:
        agrega("Módulo fuera de catálogo", d["Cod"].isin(nuevos.keys()), "Modulo",
               lambda x: d.loc[d["Modulo"] == x["Modulo"], "Area"].iloc[0],
               "área por regla: agrégalo a MODULOS_MANUAL con su nombre y área")
    sin_mod = d["Origen módulo"] == "Sin dato"
    if sin_mod.any():
        d["_mes"] = d["Fecha"].dt.strftime("%Y-%m")
        d["_maq_mes"] = d["Maquina"] + " · " + d["_mes"]
        agrega("Eventos sin módulo", sin_mod, "_maq_mes", lambda x: SIN,
               "el export no trae nivel 2 y el histórico no los tiene")
        d.drop(columns=["_mes", "_maq_mes"], inplace=True)
    for t, n in log["tipo_sin_regla"].items():
        filas.append(("Tipo sin regla", t, t, int(n), None,
                      "agrega una regla en TIPO_REGLAS"))
    agrega("Máquina sin identificar", d["Maquina"] == SIN, "Unidad", lambda x: SIN,
           "la unidad no empieza con un código de máquina")
    return pd.DataFrame(filas, columns=["Tema", "Valor", "Asignado", "Eventos", "Horas",
                                        "Qué hacer"])


def tipo_de_cambio(original, estandar, es_l4):
    partes, t = [], original
    if es_l4 and QUITAR_PREFIJO_L4 and _quitar_prefijo(original) != original:
        partes.append("prefijo de tipo quitado")
        t = _quitar_prefijo(original)
    if clave(t) != clave(estandar):
        partes.append("sinónimo (TEXTOS_MANUAL)")
    elif capitalizar(t) != estandar:
        partes.append("variante unificada")
    return " + ".join(partes) or "mayúscula inicial"


def hoja_unificaciones(d, log):
    filas = []
    for campo, col_o, col_s in (("Sistema (L3)", "L3 original", "Sistema (L3)"),
                                ("Detalle (L4)", "L4 original", "Detalle (L4)")):
        g = d[d[col_o] != d[col_s]].groupby([col_s, col_o]).size().reset_index(name="n")
        for _, x in g.iterrows():
            filas.append((campo, x[col_s], x[col_o],
                          tipo_de_cambio(x[col_o], x[col_s], campo.startswith("Detalle")),
                          int(x["n"])))
    out = pd.DataFrame(filas, columns=["Campo", "Texto estándar", "Texto original",
                                       "Cambio", "Eventos"])
    trivial = (out["Cambio"] == "mayúscula inicial").astype(int)
    return (out.assign(_t=trivial).sort_values(["Campo", "_t", "Texto estándar", "Eventos"],
                                                ascending=[True, True, True, False])
            .drop(columns="_t").reset_index(drop=True))


def tablas_resumen(d):
    tot = d["Horas"].sum() or 1
    t1 = d.pivot_table(index="Maquina", columns="Tipo", values="Horas", aggfunc="sum",
                       fill_value=0)
    t1["Total h"] = t1.sum(axis=1)
    t1["Eventos"] = d.groupby("Maquina").size()
    t1 = t1.sort_values("Total h", ascending=False).reset_index()
    d = d.assign(Periodo=d["Fecha"].dt.strftime("%Y-%m"))
    t2 = d.pivot_table(index="Periodo", columns="Tipo", values="Horas", aggfunc="sum",
                       fill_value=0)
    t2["Total h"] = t2.sum(axis=1)
    t2["Eventos"] = d.groupby("Periodo").size()
    t2 = t2.reset_index()

    def por(col, top=None):
        g = (d.groupby(col).agg(Eventos=("Horas", "size"), Horas=("Horas", "sum"))
             .sort_values("Horas", ascending=False))
        g["% horas"] = g["Horas"] / tot
        g["% acumulado"] = g["% horas"].cumsum()
        return (g.head(top) if top else g).reset_index()

    return [("Horas por máquina y tipo", t1), ("Horas por mes", t2),
            ("Horas por área", por("Area")), ("Horas por familia", por("Familia")),
            ("Top 25 módulos por horas", por("Modulo", 25))]


# =============================================================================
# 7. SALIDA EXCEL
# =============================================================================
def escribir_excel(d, calidad, pendientes, unif, resumen, ruta):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    extra = [c for c in ["Fin", "Año", "Origen módulo", "N1 original", "N2 original",
                         "L3 original", "L4 original", "Observaciones"] if c in d.columns]
    datos = d[COLUMNAS + (extra if COLUMNAS_EXTRA else [])]
    cab = PatternFill("solid", fgColor="0B2430")
    blanca = Font(bold=True, color="FFFFFF")
    negrita = Font(bold=True, color="0B2430", size=12)

    def formatear(ws, df, fila_cab=1, anchos=None, filtro=True):
        for j, col in enumerate(df.columns, start=1):
            celda = ws.cell(row=fila_cab, column=j)
            celda.fill, celda.font = cab, blanca
            celda.alignment = Alignment(vertical="center")
            muestra = [len(str(col))] + [len(str(v)) for v in df[col].head(400)]
            ancho = (anchos or {}).get(
                col, min(60, max(9, int(np.percentile(muestra, 95)) + 2)))
            ws.column_dimensions[get_column_letter(j)].width = ancho
        if filtro and len(df):
            ws.auto_filter.ref = (f"A{fila_cab}:{get_column_letter(len(df.columns))}"
                                  f"{fila_cab + len(df)}")

    def numeros(ws, df, fmt_por_col, fila_ini=2):
        for col, fmt in fmt_por_col.items():
            if col not in df.columns:
                continue
            j = df.columns.get_loc(col) + 1
            for i in range(fila_ini, fila_ini + len(df)):
                ws.cell(row=i, column=j).number_format = fmt

    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        datos.to_excel(xw, sheet_name="Datos_Delays", index=False)
        ws = xw.sheets["Datos_Delays"]
        formatear(ws, datos, anchos={"Fecha": 17, "Fin": 17, "Horas": 9, "Minutos": 9,
                                     "Mes": 6, "Año": 7, "Observaciones": 40})
        numeros(ws, datos, {"Fecha": "dd/mm/yyyy hh:mm", "Fin": "dd/mm/yyyy hh:mm",
                            "Horas": "0.000", "Minutos": "0.0"})
        ws.freeze_panes = "A2"

        fila = 1
        for titulo_t, t in resumen:
            t.to_excel(xw, sheet_name="Resumen", index=False, startrow=fila)
            ws = xw.sheets["Resumen"]
            ws.cell(row=fila, column=1, value=titulo_t).font = negrita
            formatear(ws, t, fila_cab=fila + 1, filtro=False)
            for j, col in enumerate(t.columns, start=1):
                fmt = ("0.0%" if str(col).startswith("%") else
                       "#,##0" if col == "Eventos" else
                       "#,##0.0" if j > 1 else None)
                if fmt:
                    for i in range(fila + 2, fila + 2 + len(t)):
                        ws.cell(row=i, column=j).number_format = fmt
            fila += len(t) + 4
        ws.column_dimensions["A"].width = 42

        for nombre, df, fmts in (
                ("Calidad", calidad, {"Cantidad": "#,##0"}),
                ("Pendientes", pendientes, {"Horas": "#,##0.00", "Eventos": "#,##0"}),
                ("Unificaciones", unif, {"Eventos": "#,##0"})):
            df.to_excel(xw, sheet_name=nombre, index=False)
            ws = xw.sheets[nombre]
            formatear(ws, df, anchos={"Detalle": 90, "Qué hacer": 70})
            numeros(ws, df, fmts)
            ws.freeze_panes = "A2"


# =============================================================================
# 8. MAIN
# =============================================================================
def main():
    entrada = _ruta(sys.argv[1] if len(sys.argv) > 1 else ENTRADA)
    salida = sys.argv[2] if len(sys.argv) > 2 else SALIDA
    print("=" * 70)
    print("  ESTANDARIZADOR DE DELAYS  ·  v1")
    print("=" * 70)
    if not os.path.exists(entrada):
        print(f"\n  ERROR: no encuentro el archivo:\n  {entrada}")
        print("  Ajusta ENTRADA al inicio del script o pásalo como argumento.\n")
        return

    avisos = []
    cat = Catalogo()
    cargar_historico(cat, avisos)
    cat.aplicar_manuales()
    print(f"\n  [1/6] Catálogo: {len(cat.mod)} módulos · {len(cat.fam)} sistemas con "
          f"familia · {len(cat.eventos):,} eventos en el histórico")
    for f in cat.fuentes:
        print(f"        {f}")

    d, log = estandarizar(entrada, cat, avisos)
    print(f"  [2/6] {log['leidas']:,} filas leídas · {log['duplicados']:,} duplicadas · "
          f"{log['dur_calculada']:,} con horas calculadas · "
          f"{log['maq_rescatada']:,} máquinas rescatadas")
    print(f"  [3/6] {len(d):,} filas tras filtrar tipo "
          f"({', '.join(TIPOS_INCLUIR) if TIPOS_INCLUIR else 'todos'}) y fechas "
          f"({FECHA_DESDE or 'inicio'} a {FECHA_HASTA or 'fin'})")
    if d.empty:
        print("\n  No quedó ninguna fila. Revisa TIPOS_INCLUIR, FECHA_DESDE y FECHA_HASTA.")
        print("  Tipos encontrados: " + ", ".join(
            f"{t} ({int(k):,})" for t, k in log["por_tipo"].items()) + "\n")
        return

    om = d["Origen módulo"].value_counts()
    if not log["usa_n2"]:
        avisos.insert(0, f"El módulo no viene en el export: {log['motivo_n2']}. "
                         f"Pide el export con Reason_Level2. Mientras tanto se recupera "
                         f"del histórico ({int(om.get('Histórico', 0)):,} eventos) y el "
                         f"resto queda '{SIN}' ({int(om.get('Sin dato', 0)):,}).")
    print(f"  [4/6] Módulo: {int(om.get('Nivel 2', 0)):,} del nivel 2 · "
          f"{int(om.get('Histórico', 0)):,} del histórico · "
          f"{int(om.get('Sin dato', 0)):,} sin dato")

    calidad = hoja_calidad(d, log, cat, avisos)
    pendientes = hoja_pendientes(d, log, cat)
    unif = hoja_unificaciones(d, log)
    resumen = tablas_resumen(d)
    print(f"  [5/6] {d['Sistema (L3)'].nunique()} sistemas L3 · "
          f"{d['Detalle (L4)'].nunique()} detalles L4 · "
          f"{len(pendientes)} pendientes por revisar")

    try:
        escribir_excel(d, calidad, pendientes, unif, resumen, salida)
    except PermissionError:
        base, ext = os.path.splitext(salida)
        salida = f"{base}_{datetime.now():%Y%m%d_%H%M%S}{ext}"
        print("        El archivo de salida está abierto; escribo una copia nueva.")
        escribir_excel(d, calidad, pendientes, unif, resumen, salida)
    print(f"  [6/6] {salida} ({os.path.getsize(salida) / 1024:,.0f} KB)")

    print("\n" + "=" * 70)
    print(f"  {len(d):,} {'averías' if TIPOS_INCLUIR else 'paradas'} · "
          f"{d['Horas'].sum():,.0f} h · "
          f"{d['Maquina'].nunique()} máquinas · {d.loc[d['Cod'] != SIN, 'Cod'].nunique()} "
          f"módulos")
    for a in avisos:
        print(f"  AVISO: {a}")
    print("=" * 70)

    if EJECUTAR_REPORTE:
        rep = _ruta(REPORTE_PY)
        if os.path.exists(rep):
            print()
            subprocess.run([sys.executable, rep, salida], check=False)
        else:
            print(f"\n  No encuentro {REPORTE_PY}; corre: python reporte.py {salida}")
    else:
        print(f"\n  Siguiente paso:  python reporte.py {salida}")


if __name__ == "__main__":
    main()
